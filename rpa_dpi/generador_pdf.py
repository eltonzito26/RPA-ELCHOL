import os
import json
import docx
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client

def analizar_nombre_imagen(nombre_archivo):
    nombre_sin_extension = os.path.splitext(nombre_archivo)[0]
    texto_minusculas = nombre_sin_extension.lower()

    cara = ""
    if "anverso" in texto_minusculas:
        cara = "anverso"
    elif "reverso" in texto_minusculas:
        cara = "reverso"
    else:
        cara = "desconocido"

    nombre_limpio = texto_minusculas.replace("anverso", "").replace("reverso", "")
    nombre_limpio = nombre_limpio.replace("-", " ").replace("_", " ")
    nombre_limpio = " ".join(nombre_limpio.split())
    nombre_final = nombre_limpio.title()

    return nombre_final, cara

def generar_pdfs_desde_carpeta(carpeta_entrada, carpeta_salida):
    personas = {}
    pdfs_creados = []
    incompletos = []

    # 1. Leer todos los archivos de la carpeta de entrada
    for archivo in os.listdir(carpeta_entrada):
        if archivo.lower().endswith(('.png', '.jpg', '.jpeg')):
            nombre_persona, cara = analizar_nombre_imagen(archivo)

            if nombre_persona not in personas:
                personas[nombre_persona] = {"anverso": None, "reverso": None}

            ruta_completa = os.path.join(carpeta_entrada, archivo)
            if cara == "anverso":
                personas[nombre_persona]["anverso"] = ruta_completa
            elif cara == "reverso":
                personas[nombre_persona]["reverso"] = ruta_completa

    # 2. Revisar quién tiene las dos caras y armar el PDF
    for persona, imagenes in personas.items():
        foto_anverso = imagenes["anverso"]
        foto_reverso = imagenes["reverso"]

        if foto_anverso is not None and foto_reverso is not None:
            documento = docx.Document()
            
            seccion = documento.sections[0]
            seccion.left_margin = Cm(2)
            seccion.right_margin = Cm(2)
            seccion.top_margin = Cm(2)
            seccion.bottom_margin = Cm(2)

            parrafo_anverso = documento.add_paragraph()
            parrafo_anverso.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_anverso = parrafo_anverso.add_run()
            run_anverso.add_picture(foto_anverso, width=Cm(15))

            parrafo_reverso = documento.add_paragraph()
            parrafo_reverso.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_reverso = parrafo_reverso.add_run()
            run_reverso.add_picture(foto_reverso, width=Cm(15))

            ruta_word = os.path.join(carpeta_salida, f"{persona}.docx")
            documento.save(ruta_word)

            ruta_pdf = os.path.join(carpeta_salida, f"{persona}.pdf")
            word_app = win32com.client.Dispatch("Word.Application")
            word_app.Visible = False
            
            try:
                doc_abierto = word_app.Documents.Open(ruta_word)
                doc_abierto.SaveAs(ruta_pdf, FileFormat=17)
                doc_abierto.Close()
            except Exception as error:
                print(f"Ocurrió un error con {persona}: {error}")
            finally:
                word_app.Quit()

            if os.path.exists(ruta_word):
                os.remove(ruta_word)

            pdfs_creados.append({
                "persona": persona,
                "ruta_pdf": ruta_pdf
            })

        else:
            cara_que_falta = ""
            if foto_anverso is None:
                cara_que_falta = "ANVERSO"
            elif foto_reverso is None:
                cara_que_falta = "REVERSO"

            incompletos.append({
                "persona": persona,
                "cara_faltante": cara_que_falta
            })

    # 3. Armar la respuesta final en formato JSON para que UiPath la entienda
    resultado = {
        "ok": True,
        "pdfs_creados": pdfs_creados,
        "incompletos": incompletos
    }
    
    return json.dumps(resultado, indent=4)
