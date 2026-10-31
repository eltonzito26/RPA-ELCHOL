"""Pruebas de rpa_dpi.listado con un libro sintético que imita la estructura real.

Si se define CADER_XLSX con la ruta del libro real, se ejecuta además una
prueba de integración (solo cuentas; no imprime datos personales).
"""
import os
import tempfile
import unittest
from pathlib import Path

import openpyxl

from rpa_dpi.listado import cargar_listado
from rpa_dpi.tests.test_cui import construir_cui

SEMILLAS = ["31580330", "31580331", "31580332", "31580333", "31580334", "31580335"]


def hacer_libro(ruta, filas, titulo_bloque=None, bloque=None, resumen=None):
    """filas = [(nombre, cui_celda, estado)]. Encabezado en fila 4, datos desde la 5."""
    wb = openpyxl.Workbook()
    wb.active.title = "Panel Interactivo"
    wb.create_sheet("Reporte General")
    ws = wb.create_sheet("Comunidad Uno")
    ws.cell(4, 1, "No.")
    for i, (n, c, e) in enumerate(filas, start=1):
        ws.cell(4 + i, 1, i); ws.cell(4 + i, 2, n); ws.cell(4 + i, 3, c); ws.cell(4 + i, 4, e)
    va = sum(1 for _, _, e in filas if e == "VA")
    res = resumen or (len(filas), va, len(filas) - va)
    for k, (et, v) in enumerate(zip(("Total integrantes", "DPI escaneado (Van)", "Pendientes (Faltan)"), res)):
        ws.cell(5 + k, 6, et); ws.cell(5 + k, 7, v)
    fin = 4 + len(filas)
    faltan = [f for f in filas if f[2] == "FALTA"]
    bloque = faltan if bloque is None else bloque
    ws.cell(fin + 2, 1, f"PERSONAS QUE FALTAN POR PROCESAR DPI ({titulo_bloque if titulo_bloque is not None else len(bloque)})")
    ws.cell(fin + 3, 1, "No.")
    for i, (n, c, _) in enumerate(bloque, start=1):
        ws.cell(fin + 3 + i, 1, i); ws.cell(fin + 3 + i, 2, n); ws.cell(fin + 3 + i, 3, c)
    wb.save(ruta)


class TestListadoSintetico(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.ruta = Path(self.dir.name) / "libro.xlsx"
        self.v = [construir_cui(s) for s in SEMILLAS]

    def tearDown(self):
        self.dir.cleanup()

    def tipos(self, lst):
        return sorted(p.tipo for p in lst.problemas)

    def test_libro_limpio(self):
        filas = [("A", self.v[0], "VA"), ("B", self.v[1], "FALTA"), ("C", self.v[2], "FALTA")]
        hacer_libro(self.ruta, filas)
        lst = cargar_listado(self.ruta)
        self.assertEqual(self.tipos(lst), [])
        self.assertEqual(lst.conteos()["Comunidad Uno"], {"total": 3, "VA": 1, "FALTA": 2})

    def test_busqueda_exacta(self):
        hacer_libro(self.ruta, [("A", self.v[0], "VA"), ("B", self.v[1], "FALTA")])
        lst = cargar_listado(self.ruta)
        estado, regs = lst.buscar(self.v[1])
        self.assertEqual((estado, regs[0].estado, regs[0].no), ("UNICA", "FALTA", 2))
        self.assertEqual(lst.buscar(self.v[4])[0], "SIN_COINCIDENCIA")
        self.assertEqual(lst.buscar("123")[0], "CUI_INVALIDO")

    def test_cui_con_sufijo_decimal_se_normaliza_y_se_reporta(self):
        hacer_libro(self.ruta, [("A", self.v[0] + ",00", "VA")])
        lst = cargar_listado(self.ruta)
        self.assertIn("CUI_CORREGIDO", self.tipos(lst))
        self.assertEqual(lst.buscar(self.v[0])[0], "UNICA")

    def test_duplicados_y_mod11(self):
        malo = self.v[2][:8] + str((int(self.v[2][8]) + 1) % 10) + self.v[2][9:]
        hacer_libro(self.ruta, [("A", self.v[0], "FALTA"), ("B", self.v[0], "FALTA"), ("C", malo, "FALTA")])
        lst = cargar_listado(self.ruta)
        t = self.tipos(lst)
        self.assertEqual(t.count("CUI_DUPLICADO"), 2)
        self.assertEqual(t.count("CUI_MOD11"), 1)
        self.assertEqual(lst.buscar(self.v[0])[0], "MULTIPLE")

    def test_mismo_prefijo_con_municipio_distinto(self):
        otro_muni = self.v[0][:9] + "1503"
        hacer_libro(self.ruta, [("A", self.v[0], "FALTA"), ("A2", otro_muni, "FALTA")])
        t = self.tipos(cargar_listado(self.ruta))
        self.assertEqual(t.count("CUI_PREFIJO_REPETIDO"), 2)
        self.assertNotIn("CUI_DUPLICADO", t)

    def test_bloque_y_resumen_desfasados(self):
        filas = [("A", self.v[0], "VA"), ("B", self.v[1], "FALTA"), ("C", self.v[2], "FALTA")]
        hacer_libro(self.ruta, filas, titulo_bloque=5, bloque=[filas[1], filas[0]], resumen=(3, 2, 1))
        t = self.tipos(cargar_listado(self.ruta))
        self.assertIn("BLOQUE_TITULO_DESFASADO", t)
        self.assertIn("BLOQUE_DISTINTO_DE_ESTADO", t)
        self.assertEqual(t.count("RESUMEN_DESFASADO"), 2)

    def test_el_informe_no_contiene_nombres_ni_cui(self):
        hacer_libro(self.ruta, [("NOMBRE SECRETO", self.v[0], "FALTA"), ("OTRO", self.v[0], "FALTA")])
        for p in cargar_listado(self.ruta).problemas:
            texto = f"{p.tipo}{p.comunidad}{p.detalle}"
            self.assertNotIn("SECRETO", texto)
            self.assertNotIn(self.v[0], texto)


@unittest.skipUnless(os.environ.get("CADER_XLSX"), "defina CADER_XLSX para la prueba con el libro real")
class TestLibroReal(unittest.TestCase):
    def test_totales_coinciden_con_reporte_general(self):
        lst = cargar_listado(os.environ["CADER_XLSX"])
        va = sum(1 for r in lst.registros if r.estado == "VA")
        fa = sum(1 for r in lst.registros if r.estado == "FALTA")
        self.assertEqual((len(lst.registros), va, fa), (638, 329, 309))
        self.assertEqual(len(lst.conteos()), 22)


if __name__ == "__main__":
    unittest.main()
